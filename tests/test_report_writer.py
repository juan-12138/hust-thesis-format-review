import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def sample_review(source: Path) -> dict:
    findings = [
        {
            "finding_id": "PAGE-001-0001",
            "rule_id": "PAGE-001",
            "category": "页面与分节",
            "object_type": "节",
            "severity": "ERROR",
            "problem": "纸张尺寸不符合规范",
            "actual": "宽 12240 twips，高 15840 twips",
            "expected": "A4 纵向",
            "source": "模板正文明确要求",
            "suggestion": "改为 A4 纵向",
            "location": {"scope": "section", "section_index": 0, "label": "第1节", "snippet": ""},
            "anchor": None,
            "evidence": {},
        },
        {
            "finding_id": "FONT-002-0002",
            "rule_id": "FONT-002",
            "category": "字体",
            "object_type": "正文文本",
            "severity": "WARNING",
            "problem": "正文文本颜色异常",
            "actual": "#595959",
            "expected": "#000000",
            "source": "模板正文明确要求",
            "suggestion": "改为黑色",
            "location": {"scope": "paragraph", "section_index": 1, "paragraph_index": 8, "label": "正文第9段", "snippet": "这是用于定位问题的正文片段"},
            "anchor": {"part_name": "word/document.xml", "paragraph_index": 8},
            "evidence": {},
        },
        {
            "finding_id": "PARA-003-0003",
            "rule_id": "PARA-003",
            "category": "段落",
            "object_type": "分页符",
            "severity": "MANUAL_REVIEW",
            "problem": "人工分页可能造成大面积留白",
            "actual": "检测到手工分页符",
            "expected": "结合最终分页复核",
            "source": "模板建议",
            "suggestion": "在 Word 中检查分页效果",
            "location": {"scope": "paragraph", "section_index": 1, "paragraph_index": 15, "label": "正文第16段", "snippet": "下一章内容"},
            "anchor": {"part_name": "word/document.xml", "paragraph_index": 15},
            "evidence": {},
        },
        {
            "finding_id": "HF-001-0004",
            "rule_id": "HF-001",
            "category": "页眉页脚",
            "object_type": "页眉",
            "severity": "ERROR",
            "problem": "页眉内容不符合规范",
            "actual": "错误页眉",
            "expected": "华中科技大学硕士学位论文",
            "source": "模板标准示例",
            "suggestion": "按模板设置页眉",
            "location": {"scope": "header", "section_index": 1, "label": "第2节默认页眉", "snippet": "错误页眉"},
            "anchor": {"part_name": "word/header1.xml", "paragraph_index": 0},
            "evidence": {},
        },
    ]
    return {
        "schema_version": "1.0.0-phase4",
        "source": {"path": str(source), "filename": source.name, "sha256": "ABCDEF", "size_bytes": 1234},
        "created_at_utc": "2026-09-14T12:00:00Z",
        "rule_set": {"rule_set_id": "hust-master-thesis", "schema_version": "1.0"},
        "coverage": {"implemented_rule_count": 19, "total_rule_count": 76},
        "statistics": {
            "total_findings": 4,
            "findings_by_severity": {"ERROR": 2, "WARNING": 1, "MANUAL_REVIEW": 1},
        },
        "findings": findings,
        "limitations": ["OOXML 不能可靠给出最终页码。"],
    }


def sample_comments() -> dict:
    return {
        "comment_count": 1,
        "commented_finding_count": 1,
        "uncommented_finding_count": 3,
        "findings": [
            {"finding_id": "PAGE-001-0001", "rule_id": "PAGE-001", "commented": False, "comment_id": None, "skip_reason": "missing_anchor"},
            {"finding_id": "FONT-002-0002", "rule_id": "FONT-002", "commented": True, "comment_id": 3, "skip_reason": None},
            {"finding_id": "PARA-003-0003", "rule_id": "PARA-003", "commented": False, "comment_id": None, "skip_reason": "manual_review"},
            {"finding_id": "HF-001-0004", "rule_id": "HF-001", "commented": False, "comment_id": None, "skip_reason": "unsupported_story_part"},
        ],
    }


class ReportWriterTests(unittest.TestCase):
    def test_report_contains_required_sections_all_findings_and_comment_status(self):
        from report_writer import default_report_path, write_report

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "论文初稿.docx"
            source.write_bytes(b"placeholder")
            output = default_report_path(source)
            manifest = write_report(
                source,
                sample_review(source),
                sample_comments(),
                output=output,
                created_at="2026-09-14 20:30:00",
                template_name="理工科-硕士-华中科技大学学位论文参考模板.docx",
            )
            document = Document(output)

        self.assertEqual(output.name, "论文初稿_格式审查报告.docx")
        self.assertEqual(manifest["finding_count"], 4)
        self.assertEqual(manifest["commented_finding_count"], 1)
        self.assertEqual(manifest["uncommented_finding_count"], 3)
        headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(
            headings,
            [
                "1 审查概况",
                "2 严重问题",
                "3 页面与版式问题",
                "4 字体与字符格式问题",
                "5 段落问题",
                "6 标题与编号问题",
                "7 图问题",
                "8 表问题",
                "9 公式问题",
                "10 目录问题",
                "11 页眉页脚与页码问题",
                "12 脚注、参考文献与附录问题",
                "13 无法通过 Word 批注定位的问题",
                "14 需要人工复核的问题",
                "15 审查统计",
            ],
        )
        all_text = "\n".join(p.text for p in document.paragraphs)
        all_text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        for finding in sample_review(source)["findings"]:
            self.assertIn(finding["finding_id"], all_text)
            self.assertIn(finding["problem"], all_text)
        self.assertIn("是（批注 3）", all_text)
        self.assertIn("Word XML 无法可靠获得最终页码", all_text)
        self.assertIn("当前规则覆盖中尚未包含此类专项检查", all_text)

    def test_report_has_landscape_a4_repeating_headers_borders_and_page_field(self):
        from report_writer import write_report

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "论文.docx"
            source.write_bytes(b"placeholder")
            output = temp / "报告.docx"
            manifest = write_report(source, sample_review(source), sample_comments(), output=output)
            with ZipFile(output) as package:
                self.assertIsNone(package.testzip())
                for name in package.namelist():
                    if name.endswith(".xml") or name.endswith(".rels"):
                        etree.fromstring(package.read(name))
                document_xml = etree.fromstring(package.read("word/document.xml"))
                footer_names = [name for name in package.namelist() if name.startswith("word/footer") and name.endswith(".xml")]
                footer_text = " ".join(package.read(name).decode("utf-8") for name in footer_names)

        page_size = document_xml.xpath(".//w:sectPr/w:pgSz", namespaces=NS)[0]
        self.assertEqual(page_size.get(f"{{{W_NS}}}orient"), "landscape")
        self.assertGreater(int(page_size.get(f"{{{W_NS}}}w")), int(page_size.get(f"{{{W_NS}}}h")))
        tables = document_xml.xpath(".//w:tbl", namespaces=NS)
        repeated_headers = document_xml.xpath(".//w:tbl/w:tr[1]/w:trPr/w:tblHeader", namespaces=NS)
        self.assertEqual(len(repeated_headers), len(tables))
        self.assertGreater(len(tables), 8)
        border_colors = document_xml.xpath(".//w:tblBorders/*/@w:color", namespaces=NS)
        self.assertTrue(border_colors)
        self.assertEqual(set(border_colors), {"D9D9D9"})
        header_fills = document_xml.xpath(".//w:tbl/w:tr[1]/w:tc/w:tcPr/w:shd/@w:fill", namespaces=NS)
        self.assertTrue(header_fills)
        self.assertEqual(set(header_fills), {"1F4E78"})
        self.assertIn('w:instr=" PAGE "', footer_text)
        self.assertEqual(manifest["findings_by_severity"], {"ERROR": 2, "WARNING": 1, "MANUAL_REVIEW": 1})
        self.assertEqual(manifest["manual_review_count"], 1)
        self.assertEqual(manifest["section_count"], 15)

    def test_duplicate_finding_ids_are_rejected_before_writing(self):
        from report_writer import write_report

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "论文.docx"
            source.write_bytes(b"placeholder")
            output = temp / "报告.docx"
            review = sample_review(source)
            review["findings"].append(dict(review["findings"][0]))
            with self.assertRaisesRegex(ValueError, "唯一"):
                write_report(source, review, sample_comments(), output=output)
            self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
