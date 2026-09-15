#!/usr/bin/env python3
"""Create the Phase 7 HUST thesis-format review report as a DOCX."""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


TITLE = "华中科技大学硕士学位论文格式审查报告"
DEFAULT_TEMPLATE_NAME = "华中科技大学硕士学位论文参考模板.docx"
SEVERITY_ZH = {"ERROR": "错误", "WARNING": "警告", "MANUAL_REVIEW": "人工复核"}
SKIP_REASON_ZH = {
    "manual_review": "人工复核项不写批注",
    "missing_anchor": "没有可靠锚点",
    "unsupported_story_part": "锚点位于暂不支持批注的文档部件",
    "invalid_paragraph_index": "段落索引无效",
    "paragraph_not_found": "目标段落不存在",
    "snippet_mismatch": "定位摘要与目标段落不一致",
    "unsupported_severity": "严重度不适用批注",
    "not_in_manifest": "批注映射中没有该记录",
}

DETAIL_HEADERS = ("编号", "规则", "位置", "问题与对象", "当前格式", "标准格式", "严重度", "是否已批注")
DETAIL_WIDTHS_MM = (12, 18, 52, 46, 47, 47, 17, 25)

SECTION_DEFINITIONS = (
    ("3 页面与版式问题", {"PAGE", "COVER", "FRONT", "STRUCT", "RENDER"}, True),
    ("4 字体与字符格式问题", {"FONT", "LANG"}, True),
    ("5 段落问题", {"PARA"}, True),
    ("6 标题与编号问题", {"HEAD"}, True),
    ("7 图问题", {"FIG"}, False),
    ("8 表问题", {"TABLE", "TAB"}, False),
    ("9 公式问题", {"EQUATION", "EQ"}, False),
    ("10 目录问题", {"TOC", "ABS"}, False),
    ("11 页眉页脚与页码问题", {"HF"}, True),
    ("12 脚注、参考文献与附录问题", {"FOOTNOTE", "NOTE", "REF", "BIB", "APPENDIX", "APP", "OBJ", "PKG"}, False),
)


def default_report_path(source: str | Path) -> Path:
    source = Path(source)
    return source.with_name(f"{source.stem}_格式审查报告.docx")


def _prefix(finding: dict) -> str:
    return str(finding.get("rule_id") or "").split("-", 1)[0].upper()


def _font_element(run_properties, east_asia: str, latin: str) -> None:
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:cs"), latin)


def _configure_style(style, *, east_asia: str, latin: str, size: float, bold: bool = False, alignment=None) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    _font_element(style.element.get_or_add_rPr(), east_asia, latin)
    if alignment is not None:
        style.paragraph_format.alignment = alignment


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    normal = document.styles["Normal"]
    _configure_style(normal, east_asia="宋体", latin="Times New Roman", size=10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)

    title = document.styles["Title"]
    _configure_style(title, east_asia="黑体", latin="Times New Roman", size=20, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(14)

    heading = document.styles["Heading 1"]
    _configure_style(heading, east_asia="黑体", latin="Times New Roman", size=15, bold=True)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = footer.add_run("第 ")
    _set_run_font(prefix, east_asia="宋体", latin="Times New Roman", size=9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    footer._p.append(field)
    suffix = footer.add_run(" 页")
    _set_run_font(suffix, east_asia="宋体", latin="Times New Roman", size=9)


def _set_run_font(run, *, east_asia: str, latin: str, size: float, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    _font_element(run._element.get_or_add_rPr(), east_asia, latin)


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_cell_margins(cell, *, top=80, start=100, bottom=80, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _set_table_fixed_layout(table) -> None:
    properties = table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def _style_table(table, widths_mm: tuple[float, ...], *, header=True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_fixed_layout(table)
    _set_table_borders(table)
    for column, width in zip(table.columns, widths_mm):
        column.width = Mm(width)
        for cell in column.cells:
            cell.width = Mm(width)
    if header:
        _repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index == 0 and header:
                _set_cell_shading(cell, "1F4E78")
            elif row_index % 2 == 0:
                _set_cell_shading(cell, "F3F6FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in {0, 1, 6, 7} else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    _set_run_font(
                        run,
                        east_asia="宋体",
                        latin="Times New Roman",
                        size=8.5,
                        bold=(row_index == 0 and header),
                        color="FFFFFF" if row_index == 0 and header else "000000",
                    )


def _add_table(document: Document, headers: tuple[str, ...], rows: list[list[str]], widths_mm: tuple[float, ...]):
    table = document.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
    _style_table(table, widths_mm)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _comment_map(comments: dict | None) -> dict[str, dict]:
    result = {}
    for item in (comments or {}).get("findings", []):
        finding_id = item.get("finding_id")
        if finding_id:
            result[finding_id] = item
    return result


def _location_text(finding: dict) -> str:
    location = finding.get("location") or {}
    parts = []
    section_index = location.get("section_index")
    if isinstance(section_index, int):
        parts.append(f"第 {section_index + 1} 节")
    scope = location.get("scope")
    paragraph_index = location.get("paragraph_index")
    part_name = location.get("part_name") or (finding.get("anchor") or {}).get("part_name")
    if isinstance(paragraph_index, int):
        if part_name == "word/document.xml" or scope == "paragraph":
            parts.append(f"正文段落 {paragraph_index + 1}")
        elif part_name and "header" in part_name:
            parts.append(f"页眉段落 {paragraph_index + 1}")
        elif part_name and "footer" in part_name:
            parts.append(f"页脚段落 {paragraph_index + 1}")
    label = str(location.get("label") or "").strip()
    if label and not any(token in label for token in ("word/", "¶")):
        parts.append(label)
    snippet = " ".join(str(location.get("snippet") or "").split())[:40]
    if snippet:
        parts.append(f"定位：{snippet}")
    parts.append("页码：Word XML 无法可靠获得最终页码")
    return "\n".join(dict.fromkeys(parts))


def _comment_status(finding: dict, mapping: dict[str, dict]) -> str:
    item = mapping.get(finding.get("finding_id"))
    if item and item.get("commented"):
        return f"是（批注 {item.get('comment_id')}）"
    reason = (item or {}).get("skip_reason") or "not_in_manifest"
    return f"否\n{SKIP_REASON_ZH.get(reason, reason)}"


def _detail_rows(findings: list[dict], mapping: dict[str, dict]) -> list[list[str]]:
    rows = []
    for finding in findings:
        problem = f"{finding.get('object_type', '')}：{finding.get('problem', '')}"
        source = str(finding.get("source") or "").strip()
        suggestion = str(finding.get("suggestion") or "").strip()
        if source:
            problem += f"\n依据：{source}"
        if suggestion:
            problem += f"\n建议：{suggestion}"
        rows.append(
            [
                str(finding.get("finding_id") or ""),
                str(finding.get("rule_id") or ""),
                _location_text(finding),
                problem,
                str(finding.get("actual") or ""),
                str(finding.get("expected") or ""),
                SEVERITY_ZH.get(finding.get("severity"), str(finding.get("severity") or "")),
                _comment_status(finding, mapping),
            ]
        )
    return rows


def _add_detail_section(document: Document, title: str, findings: list[dict], mapping: dict[str, dict], *, covered: bool) -> None:
    document.add_heading(title, level=1)
    if findings:
        _add_table(document, DETAIL_HEADERS, _detail_rows(findings, mapping), DETAIL_WIDTHS_MM)
    elif covered:
        document.add_paragraph("已执行当前阶段相应规则，未发现问题。")
    else:
        document.add_paragraph("当前规则覆盖中尚未包含此类专项检查，不能据此判断该类格式已经合规。")


def _add_overview(document: Document, source: Path, review: dict, comments: dict | None, *, created_at: str, template_name: str) -> None:
    findings = review.get("findings") or []
    counts = Counter(finding.get("severity") for finding in findings)
    mapping = _comment_map(comments)
    commented = sum(bool(mapping.get(finding.get("finding_id"), {}).get("commented")) for finding in findings)
    coverage = review.get("coverage") or {}
    rows = [
        ["审查文件", source.name],
        ["审查时间", created_at],
        ["标准模板", template_name],
        ["规则集", str((review.get("rule_set") or {}).get("rule_set_id") or "HUST 硕士学位论文格式规则")],
        ["当前规则覆盖", f"已实现 {coverage.get('implemented_rule_count', 0)} / 共 {coverage.get('total_rule_count', 0)} 条规则"],
        ["总问题数量", str(len(findings))],
        ["ERROR", str(counts.get("ERROR", 0))],
        ["WARNING", str(counts.get("WARNING", 0))],
        ["MANUAL REVIEW", str(counts.get("MANUAL_REVIEW", 0))],
        ["批注覆盖", f"已批注 {commented} 项，未批注 {len(findings) - commented} 项"],
        ["最终页码定位", "Word XML 无法可靠获得最终页码"],
    ]
    _add_table(document, ("项目", "结果"), rows, (45, 210))
    document.add_paragraph(
        "本报告汇总当前规则引擎发现的全部问题。ERROR 表示明确违反模板要求，WARNING 表示高概率存在问题，MANUAL REVIEW 表示需要结合 Word 最终版面人工确认。"
    )


def _add_statistics(document: Document, findings: list[dict], mapping: dict[str, dict]) -> None:
    severity = Counter(finding.get("severity") for finding in findings)
    severity_rows = [
        ["错误", severity.get("ERROR", 0)],
        ["警告", severity.get("WARNING", 0)],
        ["人工复核", severity.get("MANUAL_REVIEW", 0)],
        ["合计", len(findings)],
    ]
    _add_table(document, ("严重度", "数量"), severity_rows, (80, 45))

    categories = Counter(finding.get("category") or "未分类" for finding in findings)
    category_rows = [[name, count] for name, count in sorted(categories.items(), key=lambda item: (-item[1], item[0]))]
    _add_table(document, ("问题类别", "数量"), category_rows or [["无", 0]], (120, 45))

    commented = sum(bool(mapping.get(finding.get("finding_id"), {}).get("commented")) for finding in findings)
    comment_rows = [["已批注", commented], ["未批注", len(findings) - commented], ["批注覆盖率", f"{commented / len(findings):.1%}" if findings else "0.0%"]]
    _add_table(document, ("批注状态", "数量或比例"), comment_rows, (80, 55))


def write_report(
    source: str | Path,
    review: dict,
    comments: dict | None,
    *,
    output: str | Path | None = None,
    created_at: str | None = None,
    template_name: str = DEFAULT_TEMPLATE_NAME,
) -> dict:
    """Create the Phase 7 report and return its output/statistics manifest."""
    source = Path(source).resolve()
    output = Path(output or default_report_path(source)).resolve()
    if source == output:
        raise ValueError("审查报告不能覆盖源 DOCX")
    findings = review.get("findings")
    if not isinstance(findings, list):
        raise ValueError("review 必须包含 findings 列表")
    finding_ids = [finding.get("finding_id") for finding in findings]
    if any(not value for value in finding_ids) or len(finding_ids) != len(set(finding_ids)):
        raise ValueError("每条 Finding 必须具有唯一且非空的 finding_id")
    created_at = created_at or datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
    mapping = _comment_map(comments)

    document = Document()
    _configure_document(document)
    document.core_properties.title = TITLE
    document.core_properties.subject = "华中科技大学硕士学位论文格式审查"
    document.core_properties.author = "HUST格式审查"
    document.add_paragraph(TITLE, style="Title")

    document.add_heading("1 审查概况", level=1)
    _add_overview(document, source, review, comments, created_at=created_at, template_name=template_name)

    errors = [finding for finding in findings if finding.get("severity") == "ERROR"]
    _add_detail_section(document, "2 严重问题", errors, mapping, covered=True)

    assigned_ids = set()
    implemented_ids = set((review.get("coverage") or {}).get("implemented_rule_ids") or [])
    implemented_prefixes = {_prefix({"rule_id": rule_id}) for rule_id in implemented_ids}
    for title, prefixes, covered_by_default in SECTION_DEFINITIONS:
        selected = [finding for finding in findings if _prefix(finding) in prefixes]
        assigned_ids.update(finding.get("finding_id") for finding in selected)
        covered = covered_by_default or bool(prefixes & implemented_prefixes)
        _add_detail_section(document, title, selected, mapping, covered=covered)

    unassigned = [finding for finding in findings if finding.get("finding_id") not in assigned_ids]
    if unassigned:
        document.add_paragraph("以下问题未匹配到既定专项章节，已并入第12节：")
        _add_table(document, DETAIL_HEADERS, _detail_rows(unassigned, mapping), DETAIL_WIDTHS_MM)

    uncommented = [finding for finding in findings if not mapping.get(finding.get("finding_id"), {}).get("commented")]
    _add_detail_section(document, "13 无法通过 Word 批注定位的问题", uncommented, mapping, covered=True)

    manual = [finding for finding in findings if finding.get("severity") == "MANUAL_REVIEW"]
    _add_detail_section(document, "14 需要人工复核的问题", manual, mapping, covered=True)

    document.add_heading("15 审查统计", level=1)
    _add_statistics(document, findings, mapping)
    limitations = review.get("limitations") or []
    if limitations:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("审查限制：")
        run.bold = True
        for index, limitation in enumerate(limitations):
            if index:
                paragraph.add_run("；")
            paragraph.add_run(str(limitation).rstrip("。；"))
        paragraph.add_run("。")

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=f".{output.stem}.", suffix=".docx", dir=output.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        document.save(temporary)
        os.replace(temporary, output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise

    severity_counts = Counter(finding.get("severity") for finding in findings)
    commented_count = sum(bool(mapping.get(finding.get("finding_id"), {}).get("commented")) for finding in findings)
    return {
        "schema_version": "1.0.0-phase7",
        "output": str(output),
        "finding_count": len(findings),
        "findings_by_severity": {name: severity_counts.get(name, 0) for name in ("ERROR", "WARNING", "MANUAL_REVIEW")},
        "commented_finding_count": commented_count,
        "uncommented_finding_count": len(findings) - commented_count,
        "manual_review_count": severity_counts.get("MANUAL_REVIEW", 0),
        "section_count": 15,
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="根据阶段7审查结果和批注映射生成 Word 审查报告")
    parser.add_argument("docx", type=Path, help="原论文 DOCX，用于文件名与来源信息")
    parser.add_argument("--review", type=Path, required=True, help="阶段4审查 JSON 或阶段5完整 manifest")
    parser.add_argument("--comments", type=Path, help="阶段5批注映射 JSON；省略时尝试从 --review 读取")
    parser.add_argument("--output", type=Path, help="报告 DOCX 输出路径")
    parser.add_argument("--template-name", default=DEFAULT_TEMPLATE_NAME, help="报告中显示的标准模板名称")
    parser.add_argument("--manifest", type=Path, help="可选的报告生成结果 JSON")
    args = parser.parse_args(argv)

    payload = json.loads(args.review.read_text(encoding="utf-8"))
    review = payload.get("review") if isinstance(payload, dict) and isinstance(payload.get("review"), dict) else payload
    comments = payload.get("comments") if isinstance(payload, dict) else None
    if args.comments:
        comments_payload = json.loads(args.comments.read_text(encoding="utf-8"))
        comments = comments_payload.get("comments") if isinstance(comments_payload, dict) and isinstance(comments_payload.get("comments"), dict) else comments_payload
    result = write_report(args.docx, review, comments, output=args.output, template_name=args.template_name)
    text = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.manifest:
        args.manifest.parent.mkdir(parents=True, exist_ok=True)
        args.manifest.write_text(text, encoding="utf-8")
    else:
        print(text, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
